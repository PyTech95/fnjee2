"""FNJEE standard question-bank spreadsheet format.

One fixed format that the Import Wizard always accepts (.xlsx, .xls, .csv, .tsv):

    S.No | Subject | Class | Exam | Chapter | Topic | Question Type | Difficulty |
    Question | Option A | Option B | Option C | Option D | Option E |
    Correct Answer | Solution | Question Image | Solution Image |
    Marks | Negative Marks | Year | Source | Tags | Language

Design rules
- Header names are forgiving: "Ques", "Ans", "Option 1", "Level", "Std" … all map to
  the right column (see HEADER_ALIASES). Column order does not matter.
- Every sheet in the workbook that has a header row is imported (an "Instructions"
  sheet without a Question column is skipped). The header row may be anywhere in
  the first 15 rows.
- A sheet named after a subject ("Physics", "Bio" …) supplies the default subject.
- Pictures pasted into a row (Question Image / Solution Image cell, or anywhere on
  the row) are attached to that question. Image URLs (http…) also work.
- Answers accept A / (a) / Option B / "A, C" / AC / 1-5 (option number) / the
  option text itself / True-False / a number for integer questions.
- Question type and difficulty are auto-detected when left blank.
- Maths/chemistry formulas can be written in LaTeX between $…$ (e.g. $x^2$, $H_2O$).

Pure functions only (no DB) so they are easy to unit-test and reuse from scripts.
"""
from __future__ import annotations

import base64
import csv
import io
import re
from typing import Dict, Iterable, List, Optional, Tuple

# --------------------------------------------------------------------------------------
# Column definitions (single source of truth: template builder + parser + docs)
# --------------------------------------------------------------------------------------
# (header shown in template, internal key, required?, help text, example)
STANDARD_COLUMNS: List[Tuple[str, str, str, str, str]] = [
    ("S.No", "source_number", "Optional", "Question number in your source paper. Only for your reference.", "1"),
    ("Subject", "subject", "Required", "Physics, Chemistry, Mathematics or Biology.", "Physics"),
    ("Class", "student_class", "Optional", "11, 12 or Dropper.", "11"),
    ("Exam", "exam", "Optional", "JEE Main, JEE Advanced, NEET, Board, Olympiad …", "JEE Main"),
    ("Chapter", "chapter", "Recommended", "Chapter name. Students filter tests by this.", "Laws of Motion"),
    ("Topic", "topic", "Optional", "Sub-topic inside the chapter.", "Friction"),
    ("Question Type", "type", "Optional",
     "MCQ Single, MCQ Multiple, Integer, True/False, Assertion-Reason, Match, Subjective. "
     "Leave blank to auto-detect.", "MCQ Single"),
    ("Difficulty", "difficulty", "Optional", "Easy, Medium or Hard. Blank = Medium.", "Medium"),
    ("Question", "text", "Required",
     "Full question text. Formulas in LaTeX between $ signs, e.g. $v = u + at$, $H_2SO_4$.",
     "A block of mass $2\\,kg$ rests on a rough surface ($\\mu = 0.5$). Find the friction force."),
    ("Option A", "opt_a", "MCQ", "First option.", "$5\\,N$"),
    ("Option B", "opt_b", "MCQ", "Second option.", "$10\\,N$"),
    ("Option C", "opt_c", "MCQ", "Third option.", "$0\\,N$"),
    ("Option D", "opt_d", "MCQ", "Fourth option.", "$20\\,N$"),
    ("Option E", "opt_e", "Optional", "Fifth option (only if the question has 5).", ""),
    ("Correct Answer", "correct", "Required",
     "MCQ: A / B / C / D. Multiple correct: A,C. Integer: the number (e.g. 42 or 2.5). "
     "True/False: True or False. Subjective: leave blank and write the Solution.", "C"),
    ("Solution", "explanation", "Recommended", "Explanation shown to students after the test.",
     "The block is at rest and no horizontal force acts, so friction = 0."),
    ("Question Image", "image_url", "Optional",
     "Paste a picture into this cell OR put an image link (https://…).", ""),
    ("Solution Image", "explanation_image_url", "Optional",
     "Paste a picture into this cell OR put an image link (https://…).", ""),
    ("Marks", "marks", "Optional", "Marks for a correct answer. Blank = 4.", "4"),
    ("Negative Marks", "negative_marks", "Optional", "Marks deducted for a wrong answer. Blank = 1. Use 0 for none.", "1"),
    ("Year", "year", "Optional", "For PYQs: the exam year.", "2024"),
    ("Source", "source", "Optional", "Book / paper / file the question came from.", "DPP 05"),
    ("Tags", "tags", "Optional", "Comma separated keywords.", "pyq, important"),
    ("Language", "language", "Optional", "English or Hindi. Blank = English.", "English"),
    ("Status", "status", "Optional",
     "Approved (students can see it), Review (check before use) or Draft. Blank = Approved. "
     "Questions without an answer are always saved as Review.", "Approved"),
]

TEMPLATE_HEADERS = [c[0] for c in STANDARD_COLUMNS]

QUESTION_TYPES = {
    "mcq_single": "MCQ Single", "mcq_multi": "MCQ Multiple", "integer": "Integer",
    "true_false": "True/False", "assertion_reason": "Assertion-Reason", "match": "Match",
    "subjective": "Subjective", "image": "Image",
}
OPTION_TYPES = ("mcq_single", "mcq_multi", "true_false", "assertion_reason", "match", "image")
SUBJECTS = ("Physics", "Chemistry", "Mathematics", "Biology")


def _norm(h) -> str:
    """'Option-A *' → 'option a', 'Correct Ans.' → 'correct ans'."""
    s = str(h or "").strip().lower()
    s = s.replace("(required)", " ").replace("(optional)", " ")
    s = re.sub(r"[^a-z0-9ऀ-ॿ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


_ALIASES: Dict[str, Iterable[str]] = {
    "source_number": ["s no", "sno", "sr no", "srno", "sl no", "serial", "serial no", "q no", "qno",
                      "question no", "question number", "ques no", "no", "number", "sr", "id no"],
    "subject": ["subject", "sub", "subject name"],
    "student_class": ["class", "std", "standard", "grade", "class std"],
    "exam": ["exam", "exam type", "target", "target exam", "exam name", "category", "course"],
    "chapter": ["chapter", "chapter name", "unit", "lesson", "chap", "chapter title"],
    "topic": ["topic", "sub topic", "subtopic", "concept", "topic name"],
    "type": ["question type", "type", "q type", "qtype", "ques type", "format", "answer type"],
    "difficulty": ["difficulty", "level", "difficulty level", "toughness", "diff", "complexity"],
    "text": ["question", "question text", "questions", "ques", "q", "question statement", "problem",
             "question english", "question en", "question in english", "statement", "ques text"],
    "opt_a": ["option a", "a", "opt a", "option 1", "opt 1", "choice a", "choice 1", "optiona", "op a", "op1", "1"],
    "opt_b": ["option b", "b", "opt b", "option 2", "opt 2", "choice b", "choice 2", "optionb", "op b", "op2", "2"],
    "opt_c": ["option c", "c", "opt c", "option 3", "opt 3", "choice c", "choice 3", "optionc", "op c", "op3", "3"],
    "opt_d": ["option d", "d", "opt d", "option 4", "opt 4", "choice d", "choice 4", "optiond", "op d", "op4", "4"],
    "opt_e": ["option e", "e", "opt e", "option 5", "opt 5", "choice e", "choice 5", "optione", "op e", "op5", "5"],
    "correct": ["correct answer", "answer", "ans", "correct", "correct option", "key", "answer key",
                "right answer", "correct ans", "correct options", "answers", "solution key", "answer option"],
    "explanation": ["solution", "explanation", "explanations", "sol", "detailed solution", "solutions",
                    "answer explanation", "explaination", "reason"],
    "hint": ["hint", "hints"],
    "image_url": ["question image", "image", "image url", "diagram", "figure", "question diagram",
                  "picture", "question figure", "img", "image link", "question image url"],
    "explanation_image_url": ["solution image", "explanation image", "solution diagram", "solution figure",
                              "solution image url", "explanation image url"],
    "marks": ["marks", "mark", "positive marks", "ve marks", "score", "max marks", "correct marks"],
    "negative_marks": ["negative marks", "negative", "neg marks", "negative mark", "ve marks negative",
                       "minus marks", "penalty", "wrong marks", "negative marking"],
    "year": ["year", "exam year", "pyq year"],
    "source": ["source", "book", "reference", "paper", "file", "source file"],
    "tags": ["tags", "tag", "keywords", "labels"],
    "language": ["language", "lang", "medium"],
    "section": ["section"],
    "status": ["status"],
    "text_hi": ["question hindi", "question in hindi", "question hi", "hindi question"],
}
HEADER_ALIASES: Dict[str, str] = {}
for _key, _names in _ALIASES.items():
    for _n in _names:
        HEADER_ALIASES.setdefault(_n, _key)
for _label, _key, *_ in STANDARD_COLUMNS:
    HEADER_ALIASES[_norm(_label)] = _key

# "+ve marks" / "-ve marks" both normalise to "ve marks" – disambiguate on the raw header.
def _header_key(raw) -> Optional[str]:
    raw_s = str(raw or "").strip().lower()
    if re.match(r"^\s*-\s*ve\b", raw_s) or raw_s.startswith("(-)"):
        return "negative_marks"
    return HEADER_ALIASES.get(_norm(raw))


_TYPE_ALIASES = {
    "mcq_single": ["mcq single", "mcq", "single", "single correct", "single correct mcq", "scq", "single choice",
                   "objective", "mcq single correct", "single correct type", "mcq_single", "sc", "one correct",
                   "multiple choice", "mcq 1", "single answer"],
    "mcq_multi": ["mcq multiple", "mcq multi", "multiple", "multi", "multiple correct", "multiple correct mcq",
                  "msq", "multi correct", "more than one correct", "one or more correct", "mcq_multi", "mcc",
                  "multiple answer", "multiple choice multiple correct"],
    "integer": ["integer", "integer type", "numerical", "numeric", "nat", "numerical value",
                "integer numerical answer", "numerical answer", "integer answer", "num", "decimal"],
    "true_false": ["true false", "true or false", "t f", "tf", "true_false", "true false type"],
    "assertion_reason": ["assertion reason", "a r", "ar", "assertion and reason", "assertion_reason",
                         "assertion reasoning", "statement type"],
    "match": ["match", "match the following", "matrix match", "match the column", "match the columns",
              "column matching", "matching"],
    "subjective": ["subjective", "descriptive", "long answer", "short answer", "theory", "subjective long answer",
                   "written", "case study"],
    "image": ["image", "image based", "image based question", "diagram based", "picture based"],
}
TYPE_LOOKUP = {a: t for t, al in _TYPE_ALIASES.items() for a in al}

_DIFF_LOOKUP = {}
for _d, _al in {"easy": ["easy", "e", "1", "low", "simple", "basic", "level 1", "l1", "beginner"],
                "medium": ["medium", "m", "2", "moderate", "average", "normal", "mid", "level 2", "l2", "intermediate"],
                "hard": ["hard", "h", "3", "difficult", "tough", "high", "advanced", "level 3", "l3", "very hard"]}.items():
    for _a in _al:
        _DIFF_LOOKUP[_a] = _d

_SUBJECT_LOOKUP = {}
for _s, _al in {"Physics": ["physics", "phy", "phys", "p"],
                "Chemistry": ["chemistry", "chem", "che", "chm", "c", "organic chemistry", "inorganic chemistry",
                              "physical chemistry"],
                "Mathematics": ["mathematics", "maths", "math", "mat", "m"],
                "Biology": ["biology", "bio", "botany", "zoology", "b"]}.items():
    for _a in _al:
        _SUBJECT_LOOKUP[_a] = _s


_EXAM_PATTERNS = [
    (re.compile(r"\bjee\s*(?:adv|advanced)\b|\badvanced\b"), "JEE Advanced"),
    (re.compile(r"\bjee\s*mains?\b|\bmains\b"), "JEE Main"),
    (re.compile(r"\bneet\b"), "NEET"),
    (re.compile(r"\bjee\b|\biit\b"), "JEE"),
    (re.compile(r"\bcbse\b|\bboards?\b"), "CBSE"),
    (re.compile(r"\bolympiad\b"), "Olympiad"),
    (re.compile(r"\bcuet\b"), "CUET"),
]


def sheet_defaults(sheet: str):
    """Tab name → (subject, exam, class) defaults, e.g. 'Physics NEET' → ('Physics', 'NEET', ''),
    'JEE Main PYQ' → ('', 'JEE Main', ''), 'Bio Class 11' → ('Biology', '', '11')."""
    low = _norm(sheet)
    subject = ""
    if low in _SUBJECT_LOOKUP:
        subject = _SUBJECT_LOOKUP[low]
    else:
        for w in low.split():
            if len(w) >= 3 and w in _SUBJECT_LOOKUP:
                subject = _SUBJECT_LOOKUP[w]
                break
    exam = ""
    for rx, name in _EXAM_PATTERNS:
        if rx.search(low):
            exam = name
            break
    m = re.search(r"\b(?:class|std)\s*(9|10|11|12)\b|\b(9|10|11|12)\s*th\b", low)
    klass = (m.group(1) or m.group(2)) if m else ""
    return subject, exam, klass


def normalize_subject(v, default: str = "") -> str:
    s = _norm(v)
    if not s:
        return default
    return _SUBJECT_LOOKUP.get(s, str(v).strip().title())


def normalize_type(v) -> str:
    s = _norm(v)
    if not s:
        return ""
    if s in TYPE_LOOKUP:
        return TYPE_LOOKUP[s]
    s2 = s.replace(" ", "_")
    return s2 if s2 in QUESTION_TYPES else ""


def normalize_difficulty(v) -> str:
    return _DIFF_LOOKUP.get(_norm(v), "")


def _cell_text(v) -> str:
    """Excel cell → clean string. 2.0 → '2', keeps text exactly (incl. LaTeX)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "True" if v else "False"
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return repr(v) if len(repr(v)) < len(str(v)) else str(v)
    if hasattr(v, "isoformat") and not isinstance(v, str):
        try:
            return v.isoformat()[:10]
        except Exception:
            return str(v)
    s = str(v).replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    return s.strip()


def _to_float(v, default: float) -> float:
    s = _cell_text(v).replace("+", "").strip()
    if not s:
        return default
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    if not m:
        return default
    try:
        return float(m.group(0))
    except Exception:
        return default


_ANSWER_NOISE = re.compile(r"\b(correct|options?|opt|ans(?:wer)?s?|is|are|and|both|key|the|choice|right)\b", re.I)


def normalize_answer(raw, qtype: str, options: List[str]) -> Tuple[List[str], Optional[str]]:
    """Return (correct_list, warning)."""
    s = _cell_text(raw)
    if not s:
        return [], None
    n_opts = len(options)
    if qtype == "integer":
        num = s.replace(",", "").strip()
        m = re.fullmatch(r"\s*(-?\d+(?:\.\d+)?)\s*", num)
        if m:
            val = m.group(1)
            if "." in val:
                val = val.rstrip("0").rstrip(".") or "0"
            return [val], None
        return [s], f"answer '{s}' is not a plain number"
    if qtype == "subjective":
        return [], None

    # True / False without options
    low = s.strip().lower().rstrip(".")
    if qtype == "true_false" or (n_opts == 0 and low in ("true", "false", "t", "f")):
        if low in ("true", "t", "yes"):
            return ["A"], None
        if low in ("false", "f", "no"):
            return ["B"], None

    # Option text typed as the answer
    for i, o in enumerate(options):
        if o and low == o.strip().lower().rstrip("."):
            return [chr(65 + i)], None

    cleaned = _ANSWER_NOISE.sub(" ", s)
    cleaned = re.sub(r"[()\[\]{}.:=\-]", " ", cleaned)
    letters: List[str] = []
    for tok in re.split(r"[\s,;/&+|]+", cleaned):
        tok = tok.strip()
        if not tok:
            continue
        if re.fullmatch(r"[A-Ea-e]+", tok) and len(tok) <= 5:
            for ch in tok.upper():
                if ch not in letters:
                    letters.append(ch)
        elif re.fullmatch(r"[1-5]", tok) and n_opts >= 2:
            ch = chr(64 + int(tok))
            if ch not in letters:
                letters.append(ch)
    if not letters:
        return [], f"could not read answer '{s}'"
    if n_opts:
        bad = [l for l in letters if ord(l) - 65 >= n_opts]
        if bad:
            return letters, f"answer {','.join(bad)} but only {n_opts} options"
    return letters, None


def _detect_type(item: dict) -> str:
    opts = item.get("options") or []
    raw_ans = _cell_text(item.get("_raw_correct"))
    if len(opts) >= 2:
        letters, _ = normalize_answer(raw_ans, "mcq_multi", opts)
        return "mcq_multi" if len(letters) > 1 else "mcq_single"
    low = raw_ans.lower()
    if low in ("true", "false", "t", "f"):
        return "true_false"
    if re.fullmatch(r"-?\d+(?:\.\d+)?", raw_ans.replace(",", "")):
        return "integer"
    return "subjective"


# --------------------------------------------------------------------------------------
# Images embedded in .xlsx
# --------------------------------------------------------------------------------------
def _image_to_data_url(raw: bytes, fmt: str = "png", max_side: int = 1400) -> Optional[str]:
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(raw))
        im.load()
        if max(im.size) > max_side:
            im.thumbnail((max_side, max_side))
        out = io.BytesIO()
        if im.mode in ("RGBA", "LA", "P"):
            im.save(out, format="PNG", optimize=True)
            mime = "image/png"
        else:
            im.convert("RGB").save(out, format="JPEG", quality=85, optimize=True)
            mime = "image/jpeg"
        return f"data:{mime};base64," + base64.b64encode(out.getvalue()).decode()
    except Exception:
        if not raw:
            return None
        mime = "image/" + ("jpeg" if fmt in ("jpg", "jpeg") else (fmt or "png"))
        return f"data:{mime};base64," + base64.b64encode(raw).decode()


def _sheet_images(ws) -> Dict[int, List[Tuple[int, str]]]:
    """{row_index(1-based): [(col_index(0-based), data_url), …]}"""
    out: Dict[int, List[Tuple[int, str]]] = {}
    for img in getattr(ws, "_images", []) or []:
        try:
            anchor = img.anchor
            frm = getattr(anchor, "_from", None)
            if frm is None:
                continue
            row = frm.row + 1
            col = frm.col
            data = img._data()
            url = _image_to_data_url(data, getattr(img, "format", "png") or "png")
            if url:
                out.setdefault(row, []).append((col, url))
        except Exception:
            continue
    return out


# --------------------------------------------------------------------------------------
# Table → questions
# --------------------------------------------------------------------------------------
def _find_header(rows: List[tuple], scan: int = 15) -> Tuple[int, Dict[int, str]]:
    """Return (header_row_index, {col_index: key}) or (-1, {})."""
    best = (-1, {})
    for r_i, row in enumerate(rows[:scan]):
        cmap: Dict[int, str] = {}
        for c_i, cell in enumerate(row or ()):
            key = _header_key(cell)
            if key and key not in cmap.values():
                cmap[c_i] = key
        if "text" in cmap.values() and len(cmap) >= 2:
            # Bare "a/b/c/d/1/2/3/4" headers only count when the row also has a question column
            if len(cmap) > len(best[1]):
                best = (r_i, cmap)
    return best


def parse_question_rows(rows: List[tuple], *, sheet: str = "Sheet", default_subject: str = "",
                        images: Optional[Dict[int, List[Tuple[int, str]]]] = None,
                        row_offset: int = 0) -> Tuple[List[dict], List[str], bool]:
    """rows = list of tuples (row 0 = first spreadsheet row). Returns (questions, warnings, had_header)."""
    images = images or {}
    warnings: List[str] = []
    out: List[dict] = []
    h_idx, cmap = _find_header(rows)
    if h_idx < 0:
        return [], [], False
    sheet_subject, sheet_exam, sheet_class = sheet_defaults(sheet)
    subj_default = sheet_subject or default_subject or "Physics"
    col_of = {v: k for k, v in cmap.items()}
    blank_streak = 0
    for r_i in range(h_idx + 1, len(rows)):
        row = rows[r_i] or ()
        excel_row = r_i + 1 + row_offset
        vals: Dict[str, object] = {}
        for c_i, key in cmap.items():
            if c_i < len(row):
                vals[key] = row[c_i]
        row_imgs = images.get(excel_row, [])
        if not any(_cell_text(v) for v in vals.values()) and not row_imgs:
            blank_streak += 1
            if blank_streak > 200:
                break
            continue
        blank_streak = 0
        text = _cell_text(vals.get("text"))
        if not text:
            if row_imgs:
                text = "Refer to the figure."
            else:
                warnings.append(f"{sheet} row {excel_row}: skipped — Question cell is empty")
                continue
        # keep letter positions even when a middle option is blank (A, '', C, D)
        opt_keys = ("opt_a", "opt_b", "opt_c", "opt_d", "opt_e")
        options = [_cell_text(vals.get(k)) for k in opt_keys]
        # pictures pasted into an option cell → inline image inside that option
        opt_cols = {col_of.get(k): i for i, k in enumerate(opt_keys) if col_of.get(k) is not None}
        for c_i, url in sorted(row_imgs):
            if c_i in opt_cols:
                oi = opt_cols[c_i]
                options[oi] = (options[oi] + " ![option](" + url + ")").strip()
        while options and not options[-1]:
            options.pop()

        item: dict = {"text": text, "options": options, "_raw_correct": vals.get("correct")}
        qtype = normalize_type(vals.get("type"))
        raw_type = _cell_text(vals.get("type"))
        if raw_type and not qtype:
            warnings.append(f"{sheet} row {excel_row}: unknown Question Type '{raw_type}' — auto-detected instead")
        if qtype == "image":
            qtype = "mcq_single" if len(options) >= 2 else "image"
        if not qtype:
            qtype = _detect_type(item)
        if qtype == "true_false" and not options:
            options = ["True", "False"]
            item["options"] = options
        if qtype in ("mcq_single", "mcq_multi", "assertion_reason", "match") and len(options) < 2:
            warnings.append(f"{sheet} row {excel_row}: {QUESTION_TYPES.get(qtype, qtype)} needs at least 2 options "
                            f"— saved for review")
        correct, warn = normalize_answer(vals.get("correct"), qtype, options)
        if qtype == "mcq_single" and len(correct) > 1:
            qtype = "mcq_multi"
        if warn:
            warnings.append(f"{sheet} row {excel_row}: {warn}")
        item.pop("_raw_correct", None)
        item["type"] = qtype
        item["correct"] = correct
        item["subject"] = normalize_subject(vals.get("subject"), subj_default)
        item["difficulty"] = normalize_difficulty(vals.get("difficulty")) or "medium"
        item["marks"] = _to_float(vals.get("marks"), 4.0)
        item["negative_marks"] = abs(_to_float(vals.get("negative_marks"), 1.0))
        for key in ("chapter", "topic", "exam", "year", "source", "section", "hint", "text_hi"):
            v = _cell_text(vals.get(key))
            if v:
                item[key] = v
        if sheet_exam and not item.get("exam"):
            item["exam"] = sheet_exam
        cls = _cell_text(vals.get("student_class")) or sheet_class
        if cls:
            m = re.search(r"\b(11|12|9|10)\b", cls)
            item["student_class"] = m.group(1) if m else cls
        lang = _cell_text(vals.get("language"))
        item["language"] = lang.title() if lang else "English"
        item["explanation"] = _cell_text(vals.get("explanation"))
        tags = _cell_text(vals.get("tags"))
        if tags:
            item["tags"] = [t.strip() for t in re.split(r"[,;|]", tags) if t.strip()]
        sn = _cell_text(vals.get("source_number"))
        if sn and re.fullmatch(r"\d{1,6}", sn):
            item["source_number"] = int(sn)
        status = _norm(vals.get("status"))
        if status in ("draft", "review", "approved", "archived"):
            item["status"] = status
        else:
            item["status"] = "approved"
        # images: URL in cell, or picture anchored on this row
        for key in ("image_url", "explanation_image_url"):
            v = _cell_text(vals.get(key))
            if v and (v.startswith("http://") or v.startswith("https://") or v.startswith("data:image")):
                item[key] = v
        if row_imgs:
            q_img_col = col_of.get("image_url")
            s_img_col = col_of.get("explanation_image_url")
            sol_cols = {col_of.get("explanation"), s_img_col} - {None}
            for c_i, url in sorted(row_imgs):
                if c_i in opt_cols:
                    continue  # already placed inside the option
                if c_i in sol_cols and not item.get("explanation_image_url"):
                    item["explanation_image_url"] = url
                elif not item.get("image_url"):
                    item["image_url"] = url
                elif c_i == q_img_col:
                    item["image_url"] = url
        item["_sheet"] = sheet
        item["_row"] = excel_row
        out.append(item)
    return out, warnings, True


# --------------------------------------------------------------------------------------
# File-level entry points
# --------------------------------------------------------------------------------------
def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            t = data.decode(enc)
            if enc == "utf-16" and "\x00" in t:
                continue
            return t
        except Exception:
            continue
    return data.decode("utf-8", errors="ignore")


def parse_delimited(data: bytes, ext: str = "csv", default_subject: str = "") -> Tuple[List[dict], List[str], bool]:
    text = _decode_text(data)
    if ext == "tsv":
        delim = "\t"
    else:
        try:
            delim = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|").delimiter
        except Exception:
            delim = ","
    rows = [tuple(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]
    return parse_question_rows(rows, sheet="CSV", default_subject=default_subject)


def parse_workbook(data: bytes, ext: str = "xlsx", default_subject: str = "") -> Tuple[List[dict], List[str], dict]:
    """Parse every sheet of an .xlsx/.xls workbook. Returns (questions, warnings, summary)."""
    questions: List[dict] = []
    warnings: List[str] = []
    summary = {"sheets": []}
    if ext == "xls":
        try:
            import xlrd  # type: ignore
        except Exception:
            return [], ["Old .xls format: please open it in Excel and 'Save As' .xlsx, then upload again."], summary
        try:
            book = xlrd.open_workbook(file_contents=data)
        except Exception as e:
            return [], [f"Could not open .xls file: {e}"], summary
        for sh in book.sheets():
            rows = [tuple(sh.row_values(i)) for i in range(sh.nrows)]
            qs, ws, had = parse_question_rows(rows, sheet=sh.name, default_subject=default_subject)
            questions += qs; warnings += ws
            summary["sheets"].append({"name": sh.name, "questions": len(qs), "has_header": had})
    else:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), data_only=True)
        except Exception as e:
            return [], [f"Could not open Excel file: {e}. Make sure it is a real .xlsx file."], summary
        for ws in wb.worksheets:
            try:
                rows = list(ws.iter_rows(values_only=True))
            except Exception as e:
                warnings.append(f"Sheet '{ws.title}' could not be read: {e}")
                continue
            imgs = _sheet_images(ws)
            qs, wns, had = parse_question_rows(rows, sheet=ws.title, default_subject=default_subject, images=imgs)
            questions += qs; warnings += wns
            summary["sheets"].append({"name": ws.title, "questions": len(qs), "has_header": had})
    if not questions:
        if not any(s.get("has_header") for s in summary["sheets"]):
            warnings.insert(0, "No 'Question' column found. Download the sample file from the Import page "
                               "and keep its header row (Subject, Question, Option A … Correct Answer).")
    summary["total"] = len(questions)
    return questions, warnings, summary


def looks_like_question_table(data: bytes, ext: str) -> bool:
    """For .csv/.tsv: True when the first rows contain a Question header (vs. free text)."""
    try:
        text = _decode_text(data[:20000])
        delim = "\t" if ext == "tsv" else None
        if delim is None:
            try:
                delim = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|").delimiter
            except Exception:
                delim = ","
        rows = [tuple(r) for r in csv.reader(io.StringIO(text), delimiter=delim)][:15]
        return _find_header(rows)[0] >= 0
    except Exception:
        return False


# --------------------------------------------------------------------------------------
# Questions → standard workbook (template + export)
# --------------------------------------------------------------------------------------
def question_to_row(q: dict) -> List[object]:
    opts = list(q.get("options") or []) + [""] * 5
    t = q.get("type") or "mcq_single"
    corr = q.get("correct") or []
    if t == "true_false":
        ans = {"A": "True", "B": "False"}.get((corr or [""])[0], ",".join(corr))
    else:
        ans = ",".join(str(c) for c in corr)
    img = q.get("image_url") or ""
    sol_img = q.get("explanation_image_url") or ""
    row = {
        "source_number": q.get("source_number") or "",
        "subject": q.get("subject", ""), "student_class": q.get("student_class", ""),
        "exam": q.get("exam", ""), "chapter": q.get("chapter", ""), "topic": q.get("topic", ""),
        "type": QUESTION_TYPES.get(t, t), "difficulty": (q.get("difficulty") or "medium").title(),
        "text": q.get("text", ""),
        "opt_a": opts[0] if t != "true_false" else "", "opt_b": opts[1] if t != "true_false" else "",
        "opt_c": opts[2], "opt_d": opts[3], "opt_e": opts[4],
        "correct": ans, "explanation": q.get("explanation", ""),
        "image_url": img if img.startswith("http") else ("[picture]" if img else ""),
        "explanation_image_url": sol_img if sol_img.startswith("http") else ("[picture]" if sol_img else ""),
        "marks": q.get("marks", 4), "negative_marks": q.get("negative_marks", 1),
        "year": q.get("year", ""), "source": q.get("source", ""),
        "tags": ", ".join(q.get("tags") or []), "language": q.get("language", "English"),
        "status": (q.get("status") or "approved").title(),
    }
    return [row[c[1]] for c in STANDARD_COLUMNS]


SAMPLE_ROWS: List[dict] = [
    {"source_number": 1, "subject": "Physics", "student_class": "11", "exam": "JEE Main", "chapter": "Laws of Motion",
     "topic": "Friction", "type": "mcq_single", "difficulty": "easy",
     "text": "A block of mass $2\\,kg$ rests on a horizontal rough surface ($\\mu = 0.5$). No horizontal force acts on it. "
             "The friction force on the block is",
     "options": ["$10\\,N$", "$5\\,N$", "$0\\,N$", "$20\\,N$"], "correct": ["C"],
     "explanation": "Static friction only opposes an applied force. With no horizontal force, friction $= 0$.",
     "marks": 4, "negative_marks": 1, "year": "", "source": "Sample", "tags": ["sample"]},
    {"source_number": 2, "subject": "Chemistry", "student_class": "12", "exam": "NEET", "chapter": "Solutions",
     "topic": "Colligative properties", "type": "mcq_multi", "difficulty": "medium",
     "text": "Which of the following are colligative properties?",
     "options": ["Osmotic pressure", "Viscosity", "Elevation in boiling point", "Surface tension"],
     "correct": ["A", "C"], "explanation": "Colligative properties depend only on the number of solute particles.",
     "marks": 4, "negative_marks": 2, "source": "Sample", "tags": ["sample"]},
    {"source_number": 3, "subject": "Mathematics", "student_class": "11", "exam": "JEE Main", "chapter": "Quadratic Equations",
     "topic": "Roots", "type": "integer", "difficulty": "medium",
     "text": "If $\\alpha$ and $\\beta$ are roots of $x^2 - 5x + 6 = 0$, find $\\alpha^2 + \\beta^2$.",
     "options": [], "correct": ["13"], "explanation": "$\\alpha^2+\\beta^2 = (\\alpha+\\beta)^2 - 2\\alpha\\beta = 25 - 12 = 13$.",
     "marks": 4, "negative_marks": 0, "source": "Sample", "tags": ["sample"]},
    {"source_number": 4, "subject": "Biology", "student_class": "11", "exam": "NEET", "chapter": "Cell: The Unit of Life",
     "topic": "Cell organelles", "type": "true_false", "difficulty": "easy",
     "text": "Mitochondria contain their own circular DNA.", "options": [], "correct": ["A"],
     "explanation": "Mitochondria have their own circular DNA and 70S ribosomes.",
     "marks": 4, "negative_marks": 1, "source": "Sample", "tags": ["sample"]},
    {"source_number": 5, "subject": "Biology", "student_class": "12", "exam": "NEET", "chapter": "Molecular Basis of Inheritance",
     "topic": "DNA replication", "type": "assertion_reason", "difficulty": "hard",
     "text": "Assertion (A): DNA replication is semi-conservative.\nReason (R): Each new DNA molecule has one parental and one newly synthesised strand.",
     "options": ["Both A and R are true and R is the correct explanation of A",
                 "Both A and R are true but R is not the correct explanation of A",
                 "A is true but R is false", "A is false but R is true"],
     "correct": ["A"], "explanation": "Meselson–Stahl experiment proved semi-conservative replication; R explains A.",
     "marks": 4, "negative_marks": 1, "source": "Sample", "tags": ["sample"]},
    {"source_number": 6, "subject": "Physics", "student_class": "12", "exam": "Board", "chapter": "Electrostatics",
     "topic": "Coulomb's law", "type": "subjective", "difficulty": "medium",
     "text": "State Coulomb's law and write it in vector form.", "options": [], "correct": [],
     "explanation": "The force between two point charges is $F = \\dfrac{1}{4\\pi\\varepsilon_0}\\dfrac{q_1 q_2}{r^2}$ along the line joining them. "
                    "Vector form: $\\vec F_{12} = \\dfrac{1}{4\\pi\\varepsilon_0}\\dfrac{q_1 q_2}{r^2}\\hat r_{12}$.",
     "marks": 3, "negative_marks": 0, "source": "Sample", "tags": ["sample"]},
]


def build_workbook(questions: Optional[List[dict]] = None, *, with_instructions: bool = True) -> bytes:
    """Build the standard workbook. questions=None → sample template with example rows."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Questions"
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    req_fill = PatternFill("solid", fgColor="B03A2E")
    thin = Side(style="thin", color="D0D7DE")
    ws.append(TEMPLATE_HEADERS)
    for i, (label, key, req, *_rest) in enumerate(STANDARD_COLUMNS, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = req_fill if req == "Required" else head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=thin)
    rows = questions if questions is not None else SAMPLE_ROWS
    for q in rows:
        ws.append(question_to_row(q))
        for c in ws[ws.max_row]:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.data_type = "s"     # '= 1' is text, not an Excel formula
    widths = {"text": 60, "explanation": 50, "opt_a": 22, "opt_b": 22, "opt_c": 22, "opt_d": 22, "opt_e": 14,
              "chapter": 24, "topic": 20, "type": 16, "correct": 14, "subject": 13, "exam": 12,
              "image_url": 18, "explanation_image_url": 18, "source": 16, "tags": 16}
    for i, (_l, key, *_r) in enumerate(STANDARD_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(key, 11)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "J2"
    ws.row_dimensions[1].height = 30
    max_r = max(ws.max_row + 5000, 5000)
    col = {k: get_column_letter(i) for i, (_l, k, *_r) in enumerate(STANDARD_COLUMNS, start=1)}

    def _dv(key, values):
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                            showErrorMessage=False)
        dv.add(f"{col[key]}2:{col[key]}{max_r}")
        ws.add_data_validation(dv)
    _dv("subject", list(SUBJECTS))
    _dv("type", ["MCQ Single", "MCQ Multiple", "Integer", "True/False", "Assertion-Reason", "Match", "Subjective"])
    _dv("difficulty", ["Easy", "Medium", "Hard"])
    _dv("student_class", ["11", "12", "Dropper"])
    _dv("exam", ["JEE Main", "JEE Advanced", "NEET", "Board", "Olympiad", "CUET"])
    _dv("language", ["English", "Hindi"])
    _dv("status", ["Approved", "Review", "Draft"])

    if with_instructions:
        ins = wb.create_sheet("Instructions")
        ins.append(["FNJEE Question Upload Format"])
        ins["A1"].font = Font(bold=True, size=14)
        ins.append(["Fill one question per row in the 'Questions' sheet. Keep the header row as it is. "
                    "Red headers are required. Delete the sample rows before uploading your own questions."])
        ins.append(["You can add more sheets (e.g. one sheet per chapter) — every sheet with a 'Question' header is imported."])
        ins.append([])
        ins.append(["Column", "Required?", "What to write", "Example"])
        for cidx in range(1, 5):
            ins.cell(row=5, column=cidx).font = Font(bold=True, color="FFFFFF")
            ins.cell(row=5, column=cidx).fill = head_fill
        for label, _key, req, help_, ex in STANDARD_COLUMNS:
            ins.append([label, req, help_, ex])
        ins.append([])
        ins.append(["Answer formats", "", "A  |  (b)  |  Option C  |  A,C (multiple correct)  |  1-4 = A-D  |  42 or 2.5 (integer)  |  True / False", ""])
        ins.append(["Formulas", "", "Write maths/chemistry in LaTeX between $ signs: $x^2$, $\\frac{a}{b}$, $\\sqrt{2}$, $H_2SO_4$, $10^{-3}$", ""])
        ins.append(["Pictures", "", "Insert → Picture → Place in cell (or paste) in the 'Question Image' / 'Solution Image' column of that row.", ""])
        ins.append(["Upload", "", "Admin → Import Wizard → choose this file (or many files at once) → Upload. "
                    "Tick 'Save directly' for large files to skip the review screen.", ""])
        ins.column_dimensions["A"].width = 20
        ins.column_dimensions["B"].width = 13
        ins.column_dimensions["C"].width = 95
        ins.column_dimensions["D"].width = 45
        for row in ins.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template() -> bytes:
    return build_workbook(None, with_instructions=True)
