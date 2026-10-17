"""Standard FNJEE Excel format: template, forgiving parser, bulk save-directly, export round trip.

Runs fully in-process against an in-memory Mongo (mongomock-motor) — no live server needed:
    cd backend && python -m pytest -n 0 -p no:cacheprovider tests/test_standard_excel_import.py
"""
import asyncio
import json
import io
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("MONGO_URL", "mongodb://localhost:27017")
os.environ.setdefault("DB_NAME", "fnjee_test")
os.environ.setdefault("JWT_SECRET", "x" * 40)
os.environ.setdefault("CORS_ORIGINS", "http://localhost:3000")

mongomock_motor = pytest.importorskip("mongomock_motor")
httpx = pytest.importorskip("httpx")

import question_sheet as qs  # noqa: E402


@pytest.fixture()
def app_client():
    import server
    server.db = mongomock_motor.AsyncMongoMockClient()["fnjee_test"]
    admin = {"id": "admin-1", "role": "admin", "email": "a@x.com", "name": "Admin"}

    async def _admin():
        return admin
    server.app.dependency_overrides[server.get_current_user] = _admin
    transport = httpx.ASGITransport(app=server.app)
    client = httpx.AsyncClient(transport=transport, base_url="http://test")
    yield server, client
    server.app.dependency_overrides.clear()


async def _wait_job(client, job_id):
    for _ in range(200):
        r = await client.get(f"/api/import/jobs/{job_id}")
        j = r.json()
        if j["status"] in ("done", "error"):
            return j
        await asyncio.sleep(0.05)
    raise AssertionError("job did not finish")


def _xlsx(rows, title="Questions"):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = title
    for r in rows:
        ws.append(r)
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def test_template_parses_back_to_six_questions():
    data = qs.build_template()
    questions, warnings, summary = qs.parse_workbook(data)
    assert summary["total"] == 6 and not warnings
    types = [q["type"] for q in questions]
    assert types == ["mcq_single", "mcq_multi", "integer", "true_false", "assertion_reason", "subjective"]
    assert questions[1]["correct"] == ["A", "C"]
    assert questions[2]["correct"] == ["13"]


def test_forgiving_headers_and_answers():
    rows = [("Sr No", "Ques", "Opt 1", "Opt 2", "Opt 3", "Opt 4", "Ans", "Level", "Sol"),
            (1, "2+2?", 3, 4, 5, 6, "(b)", "E", "four"),
            (2, "Primes?", 2, 3, 4, 6, "1, 2", "tough", ""),
            (3, "g?", None, None, None, None, 9.8, "", ""),
            (4, "Earth is flat", None, None, None, None, "False", "", ""),
            (5, "Unit of energy", "Newton", "Joule", "Watt", "Pascal", "Joule", "", "")]
    q, w, had = qs.parse_question_rows(rows, sheet="Chemistry")
    assert had and len(q) == 5
    assert [x["correct"] for x in q] == [["B"], ["A", "B"], ["9.8"], ["B"], ["B"]]
    assert [x["type"] for x in q] == ["mcq_single", "mcq_multi", "integer", "true_false", "mcq_single"]
    assert all(x["subject"] == "Chemistry" for x in q)  # from sheet name
    assert q[0]["difficulty"] == "easy" and q[1]["difficulty"] == "hard"


def test_embedded_picture_attached_to_row():
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image
    wb = Workbook(); ws = wb.active
    ws.append(qs.TEMPLATE_HEADERS)
    ws.append(qs.question_to_row({"subject": "Physics", "text": "Find current in the circuit shown.",
                                  "options": ["1 A", "2 A", "3 A", "4 A"], "correct": ["B"], "type": "mcq_single"}))
    pic = io.BytesIO(); Image.new("RGB", (40, 30), "red").save(pic, format="PNG"); pic.seek(0)
    img = XLImage(pic)
    col = qs.TEMPLATE_HEADERS.index("Question Image") + 1
    from openpyxl.utils import get_column_letter
    ws.add_image(img, f"{get_column_letter(col)}2")
    b = io.BytesIO(); wb.save(b)
    questions, _, _ = qs.parse_workbook(b.getvalue())
    assert len(questions) == 1
    assert questions[0]["image_url"].startswith("data:image/")


@pytest.mark.asyncio
async def test_template_endpoint_and_bulk_save_directly(app_client):
    server, client = app_client
    r = await client.get("/api/import/template")
    assert r.status_code == 200 and r.content[:2] == b"PK"

    rows = [qs.TEMPLATE_HEADERS]
    for i in range(1, 1201):
        rows.append(qs.question_to_row({"subject": "Physics", "chapter": "Kinematics", "type": "mcq_single",
                                        "text": f"Bulk question number {i}?", "options": ["a", "b", "c", "d"],
                                        "correct": ["A"], "explanation": "because"}))
    data = _xlsx(rows)
    files = {"file": ("bulk.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = await client.post("/api/import/start", files=files, data={"auto_commit": "true", "use_ai": "false"})
    assert r.status_code == 200, r.text
    job = await _wait_job(client, r.json()["job_id"])
    assert job["status"] == "done", job
    res = job["result"]
    assert res["count"] == 1200 and res["commit"]["inserted"] == 1200
    assert await server.db.questions.count_documents({}) == 1200

    # upload the same file again → all duplicates skipped
    r = await client.post("/api/import/start", files=files, data={"auto_commit": "true", "use_ai": "false"})
    job = await _wait_job(client, r.json()["job_id"])
    assert job["result"]["commit"]["inserted"] == 0
    assert job["result"]["commit"]["duplicates_skipped"] == 1200
    assert await server.db.questions.count_documents({}) == 1200


@pytest.mark.asyncio
async def test_review_mode_returns_all_rows_then_commit(app_client):
    server, client = app_client
    data = qs.build_template()
    files = {"file": ("FNJEE_Question_Upload_Format.xlsx", data, "application/octet-stream")}
    r = await client.post("/api/import/start", files=files, data={"use_ai": "false"})
    job = await _wait_job(client, r.json()["job_id"])
    res = job["result"]
    assert res["count"] == 6 and len(res["questions"]) == 6
    r = await client.post("/api/import/commit", json={"questions": res["questions"], "skip_duplicates": True})
    assert r.json()["inserted"] == 6
    subj = await server.db.questions.find_one({"type": "subjective"}, {"_id": 0})
    assert subj["status"] == "approved"  # has a solution

    r = await client.get("/api/questions/export")
    assert r.status_code == 200
    back, _, _ = qs.parse_workbook(r.content)
    assert len(back) == 6 and back[1]["correct"] == ["A", "C"]


@pytest.mark.asyncio
async def test_csv_with_header_uses_table_parser(app_client):
    _server, client = app_client
    csv_data = "Subject,Question,Option A,Option B,Option C,Option D,Correct Answer\n" \
               "Biology,Powerhouse of the cell?,Nucleus,Mitochondria,Ribosome,Golgi,B\n"
    files = {"file": ("q.csv", csv_data.encode(), "text/csv")}
    r = await client.post("/api/import/start", files=files, data={"use_ai": "false"})
    job = await _wait_job(client, r.json()["job_id"])
    q = job["result"]["questions"]
    assert len(q) == 1 and q[0]["correct"] == ["B"] and q[0]["subject"] == "Biology"


def test_sheet_name_defaults():
    rows = [("Question", "Option A", "Option B", "Correct Answer"), ("Pick one", "x", "y", "A")]
    q, _, _ = qs.parse_question_rows(rows, sheet="Physics NEET Class 12")
    assert q[0]["subject"] == "Physics" and q[0]["exam"] == "NEET" and q[0]["student_class"] == "12"


@pytest.mark.asyncio
async def test_pictures_saved_as_files_not_base64(app_client, tmp_path, monkeypatch):
    import media_store
    monkeypatch.setattr(media_store, "MEDIA_DIR", tmp_path)
    server, client = app_client
    from PIL import Image
    buf = io.BytesIO(); Image.new("RGB", (30, 20), "blue").save(buf, format="PNG")
    url = "data:image/png;base64," + __import__("base64").b64encode(buf.getvalue()).decode()
    q = {"subject": "Physics", "type": "mcq_single", "text": f"Look ![fig]({url}) here",
         "options": [f"![o]({url})", "b"], "correct": ["A"], "image_url": url, "explanation": "e"}
    r = await client.post("/api/import/commit", json={"questions": [q]})
    assert r.json()["inserted"] == 1
    doc = await server.db.questions.find_one({}, {"_id": 0})
    assert doc["image_url"].startswith("/api/media/") and "base64" not in str(doc)
    assert len(list(tmp_path.iterdir())) == 1          # same picture stored once
    r = await client.get(doc["image_url"])
    assert r.status_code == 200 and r.content == buf.getvalue()
    assert (await client.get("/api/media/../../etc/passwd")).status_code == 404


@pytest.mark.asyncio
async def test_same_stem_different_options_is_not_duplicate(app_client):
    _server, client = app_client
    a = {"subject": "Physics", "text": "Which of the following is correct?", "options": ["a", "b", "c", "d"], "correct": ["A"]}
    b = {"subject": "Physics", "text": "Which of the following is correct?", "options": ["w", "x", "y", "z"], "correct": ["B"]}
    r = await client.post("/api/import/commit", json={"questions": [a, b, dict(a)], "skip_duplicates": True})
    j = r.json()
    assert j["inserted"] == 2 and j["duplicates_skipped"] == 1


@pytest.mark.asyncio
async def test_picture_questions_with_same_text_are_not_duplicates(app_client, tmp_path, monkeypatch):
    import media_store, base64
    from PIL import Image
    monkeypatch.setattr(media_store, "MEDIA_DIR", tmp_path)
    _server, client = app_client
    def pic(color):
        b = io.BytesIO(); Image.new("RGB", (20, 20), color).save(b, format="PNG")
        return "data:image/png;base64," + base64.b64encode(b.getvalue()).decode()
    base = {"subject": "Mathematics", "text": "Refer to the question shown in the figure.",
            "options": ["(a)", "(b)", "(c)", "(d)"], "correct": ["A"]}
    qs_ = [dict(base, image_url=pic("red")), dict(base, image_url=pic("blue")), dict(base, image_url=pic("red"))]
    j = (await client.post("/api/import/commit", json={"questions": qs_, "skip_duplicates": True})).json()
    assert j["inserted"] == 2 and j["duplicates_skipped"] == 1
    # re-upload: both now exist in the bank (stored as /api/media/ files) → both skipped
    j = (await client.post("/api/import/commit", json={"questions": qs_[:2], "skip_duplicates": True})).json()
    assert j["inserted"] == 0 and j["duplicates_skipped"] == 2


def _png(text="Q1. Find x", size=(500, 160)):
    from PIL import Image, ImageDraw
    im = Image.new("RGB", size, "white"); ImageDraw.Draw(im).text((10, 10), text, fill="black")
    b = io.BytesIO(); im.save(b, format="PNG"); return b.getvalue()


@pytest.mark.asyncio
async def test_screenshot_import_flow(app_client, monkeypatch):
    server, client = app_client
    import ai_chat
    files = {"question_images": ("q.png", _png(), "image/png")}
    monkeypatch.setattr(server, "ai_status", lambda: {"masked_key": ""})
    r = await client.post("/api/import/screenshot/start", files=files)
    assert r.status_code == 400 and "AI Settings" in r.json()["detail"]

    monkeypatch.setattr(server, "ai_status", lambda: {"masked_key": "AIza…abcd"})
    seen = {}
    async def fake_ai(system, text, file_paths=None, max_tokens=8192):
        import pymupdf
        d = pymupdf.open(file_paths[0]); seen["pages"] = [p.get_text().split("\n")[0] for p in d]
        return json.dumps({"questions": [{"type": "mcq_single", "text": "Find $x$ if $2x=4$",
            "options": ["1", "2", "3", "4"], "correct": ["b"], "answer_source": "screenshot",
            "explanation": "$x = 2$", "has_diagram": True, "diagram": {"page": 1, "bbox": [0, 100, 500, 900]},
            "subject": "Mathematics", "chapter": "Linear Equations", "difficulty": "easy"}], "warnings": []})
    monkeypatch.setattr(ai_chat, "ai_complete", fake_ai)
    multi = [("question_images", ("q.png", _png(), "image/png")),
             ("answer_images", ("a.png", _png("Ans (b) x=2"), "image/png"))]
    r = await client.post("/api/import/screenshot/start", files=multi, data={"subject_default": "Mathematics"})
    assert r.status_code == 200, r.text
    job = await _wait_job(client, r.json()["job_id"])
    assert job["status"] == "done", job
    q = job["result"]["questions"][0]
    assert seen["pages"] == ["QUESTION 1", "ANSWER 1"]
    assert q["correct"] == ["B"] and q["status"] == "approved" and q["image_url"].startswith("data:image/png")
    r = await client.post("/api/import/commit", json={"questions": [q]})
    assert r.json()["inserted"] == 1
